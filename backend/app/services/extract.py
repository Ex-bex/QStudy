MAX_CHARS = 150_000  # предохранитель от гигантских файлов

def _clip(text: str) -> str:
    return text[:MAX_CHARS]

def extract_pdf(path: str) -> str:
    import fitz  # pymupdf
    parts = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return _clip("\n\n".join(parts))

def extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return _clip("\n".join(parts))

def extract_xlsx(path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"### Лист: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
            if len(parts) > 5000:
                parts.append("...(файл обрезан)")
                break
    wb.close()
    return _clip("\n".join(parts))

def extract_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"### Слайд {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    parts.append(txt)
    return _clip("\n".join(parts))

def extract_txt(path: str) -> str:
    raw = open(path, "rb").read()
    for enc in ("utf-8", "cp1251"):  # вторая — на случай старых файлов с винды
        try:
            return _clip(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    return _clip(raw.decode("utf-8", errors="replace"))

def extract_document(path: str, file_type: str) -> str:
    if file_type == "pdf":
        return extract_pdf(path)
    if file_type == "docx":
        return extract_docx(path)
    if file_type == "xlsx":
        return extract_xlsx(path)
    if file_type == "pptx":
        return extract_pptx(path)
    if file_type == "text":
        return extract_txt(path)
    raise ValueError(f"Тип {file_type} не поддерживается локальным извлечением")