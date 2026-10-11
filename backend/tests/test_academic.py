"""Academic API regression tests use an isolated in-memory database and no AI network calls."""
import io
import asyncio
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

os.environ.setdefault("SECRET_KEY", "test-only-qstudy-secret-key-32-characters-long")

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth import get_current_user
from app.config import settings
from app.database import Base, get_db
from app.models import ChatSession, Course, CurriculumWeek, ModuleWork, StudentGroup, SyllabusDocument, User, WorkSubmission
from app.routers.academic import router
from app.routers import chat
from app.services.syllabus import SyllabusDraft, analyze_syllabus_text, extract_syllabus_bytes, seed_academic_defaults
from app.services.reference_curriculum import REFERENCE_DATA, seed_sourced_curriculum
from app.services.academic_context import MAX_ACADEMIC_CONTEXT, academic_context
from app.services import background


class AcademicApiTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        @event.listens_for(self.engine, "connect")
        def foreign_keys(connection, record):
            connection.execute("PRAGMA foreign_keys=ON")
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, expire_on_commit=False)
        with self.sessions() as db:
            seed_academic_defaults(db)
            seed_academic_defaults(db)
            self.group_id = db.query(StudentGroup).one().id
            other_group = StudentGroup(name="Другая группа")
            db.add(other_group)
            db.flush()
            self.other_group_id = other_group.id
            people = [User(email="admin@test.example", name="Администратор", hashed_password="unused", role="admin", group_id=self.group_id),
                      User(email="student@test.example", name="Студент", hashed_password="unused", role="student", group_id=self.group_id),
                      User(email="peer@test.example", name="Другой студент", hashed_password="unused", role="student", group_id=self.group_id),
                      User(email="outsider@test.example", name="Другая группа", hashed_password="unused", role="student", group_id=other_group.id)]
            db.add_all(people)
            db.commit()
            self.admin, self.student, self.peer, self.outsider = people
        self.actor = self.admin
        self.temp = tempfile.TemporaryDirectory()
        self.previous_upload = settings.UPLOAD_DIR
        self.previous_key = settings.GEMINI_API_KEY
        settings.GEMINI_API_KEY = ""
        settings.UPLOAD_DIR = self.temp.name
        app = FastAPI()
        app.include_router(router, prefix="/api")
        app.include_router(chat.router, prefix="/api")
        def database():
            with self.sessions() as db:
                yield db
        app.dependency_overrides[get_db] = database
        app.dependency_overrides[get_current_user] = lambda: self.actor
        self.client = TestClient(app)
        self.course = self.client.post(f"/api/groups/{self.group_id}/courses", json={"title": "Профессиональный русский язык"}).json()

    def tearDown(self):
        self.client.close()
        settings.UPLOAD_DIR = self.previous_upload
        settings.GEMINI_API_KEY = self.previous_key
        self.engine.dispose()
        self.temp.cleanup()

    def create_work(self):
        self.actor = self.admin
        result = self.client.post(f"/api/courses/{self.course['id']}/works", json={"title": "Модуль 1", "kind": "module", "week_number": 2, "max_score": 20})
        self.assertEqual(result.status_code, 201, result.text)
        return result.json()

    def upload(self):
        self.actor = self.admin
        document = Document()
        document.add_paragraph("Профессиональный русский язык. Неделя 1. Тема: профессиональное общение. Задание: составить диалог.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Неделя 2"
        table.cell(0, 1).text = "Терминология специальности"
        content = io.BytesIO()
        document.save(content)
        result = self.client.post("/api/syllabi/upload", data={"course_id": self.course["id"]}, files={"file": ("../../syllabus.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(result.status_code, 201, result.text)
        self.assertNotIn("filepath", result.json())
        self.assertIn("Терминология специальности", result.json()["extracted_text"])
        return result.json()

    def test_students_can_read_but_cannot_mutate_academic_records(self):
        work = self.create_work()
        week = self.client.post(f"/api/courses/{self.course['id']}/weeks", json={"week_number": 1, "title": "Тема"}).json()
        self.actor = self.student
        self.assertEqual(self.client.get("/api/groups").status_code, 200)
        self.assertEqual(self.client.get(f"/api/courses/{self.course['id']}").status_code, 200)
        requests = [("post", "/api/groups", {"name": "Подмена"}),
                    ("put", f"/api/groups/{self.group_id}", {"name": "Подмена"}),
                    ("post", f"/api/groups/{self.group_id}/courses", {"title": "Подмена"}),
                    ("put", f"/api/courses/{self.course['id']}", {"title": "Подмена"}),
                    ("put", f"/api/weeks/{week['id']}", {"title": "Подмена", "week_number": 1}),
                    ("put", f"/api/works/{work['id']}", {"title": "Подмена"})]
        for method, url, data in requests:
            self.assertEqual(getattr(self.client, method)(url, json=data).status_code, 403, url)
        self.assertEqual(self.client.delete(f"/api/groups/{self.group_id}").status_code, 403)
        self.assertEqual(self.client.get(f"/api/works/{work['id']}/submissions").status_code, 403)

    def test_schedule_validates_group_time_and_conflicts(self):
        endpoint = f"/api/groups/{self.group_id}/schedule"
        payload = {"course_id": self.course["id"], "day_of_week": 1, "start_time": "09:00", "end_time": "10:30", "room": "305", "week_pattern": "odd"}
        self.assertEqual(self.client.post(endpoint, json=payload).status_code, 201)
        self.assertEqual(self.client.post(endpoint, json=payload).status_code, 409)
        self.assertEqual(self.client.post(endpoint, json={**payload, "week_pattern": "even"}).status_code, 201)
        self.assertEqual(self.client.post(endpoint, json={**payload, "end_time": "08:00"}).status_code, 422)
        self.assertEqual(self.client.post(f"/api/groups/{self.other_group_id}/schedule", json=payload).status_code, 422)
        schedule = self.client.get(f"/api/groups/{self.group_id}").json()["schedule"]
        self.assertEqual(schedule[0]["course_title"], self.course["title"])

    def test_schedule_calendar_filter_uses_iso_week_parity_and_year_boundary(self):
        endpoint = f"/api/groups/{self.group_id}/schedule"
        for day in (1, 5):
            for pattern in ("odd", "even"):
                result = self.client.post(endpoint, json={"course_id": self.course["id"], "day_of_week": day,
                                                        "start_time": "09:00", "end_time": "10:00", "week_pattern": pattern})
                self.assertEqual(result.status_code, 201, result.text)
        self.client.post(endpoint, json={"course_id": self.course["id"], "day_of_week": 1,
                                        "start_time": "11:00", "end_time": "12:00", "week_pattern": "every"})
        for on, expected in [("2026-10-05", ["odd", "every"]), ("2026-10-12", ["even", "every"]),
                             ("2027-01-01", ["odd"]), ("2027-01-04", ["odd", "every"]), ("2026-10-10", [])]:
            result = self.client.get(endpoint, params={"on": on})
            self.assertEqual(result.status_code, 200, result.text)
            self.assertEqual([entry["week_pattern"] for entry in result.json()], expected)
        self.assertEqual(len(self.client.get(endpoint).json()), 5)
        self.assertEqual(self.client.get(endpoint, params={"on": "2026-13-40"}).status_code, 422)

    def test_work_deadline_is_normalized_and_returned_with_explicit_utc(self):
        result = self.client.post(f"/api/courses/{self.course['id']}/works", json={"title": "Срок", "due_date": "2026-10-10T15:30:00+05:00"})
        self.assertEqual(result.status_code, 201, result.text)
        self.assertEqual(result.json()["due_date"], "2026-10-10T10:30:00+00:00")

    def test_submissions_are_owned_group_scoped_and_grades_protected(self):
        work = self.create_work()
        endpoint = f"/api/works/{work['id']}/submission"
        self.actor = self.student
        self.assertEqual(self.client.put(endpoint, json={"status": "submitted", "content": ""}).status_code, 422)
        submission = self.client.put(endpoint, json={"content": "Мой ответ", "status": "submitted"}).json()
        self.assertEqual(self.client.put(endpoint, json={"content": "Ответ", "grade": 20}).status_code, 422)
        self.actor = self.peer
        self.assertIsNone(self.client.get(f"/api/works/{work['id']}").json()["my_submission"])
        self.client.delete(endpoint)
        self.actor = self.outsider
        self.assertEqual(self.client.put(endpoint, json={"content": "Чужой ответ"}).status_code, 403)
        self.actor = self.admin
        self.assertEqual(self.client.put(f"/api/submissions/{submission['id']}/grade", json={"grade": 21}).status_code, 422)
        result = self.client.put(f"/api/submissions/{submission['id']}/grade", json={"grade": 18, "feedback": "Хорошая работа"})
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(len(self.client.get(f"/api/works/{work['id']}/submissions").json()), 1)
        self.actor = self.student
        self.assertEqual(self.client.put(endpoint, json={"content": "Заменить"}).status_code, 409)
        self.assertEqual(self.client.delete(endpoint).status_code, 409)
        self.assertEqual(self.client.get(f"/api/works/{work['id']}").json()["my_submission"]["grade"], 18)

    def test_progress_is_private_and_unique(self):
        week = self.client.post(f"/api/courses/{self.course['id']}/weeks", json={"week_number": 1, "title": "Введение"}).json()
        self.assertEqual(self.client.post(f"/api/courses/{self.course['id']}/weeks", json={"week_number": 1, "title": "Дубликат"}).status_code, 409)
        self.actor = self.student
        endpoint = f"/api/weeks/{week['id']}/progress"
        self.client.put(endpoint, json={"completed": True, "note": "Личная заметка"})
        self.client.put(endpoint, json={"completed": False, "note": "Повторить"})
        self.assertEqual(self.client.get(f"/api/courses/{self.course['id']}").json()["weeks"][0]["my_progress"]["note"], "Повторить")
        self.actor = self.peer
        self.assertIsNone(self.client.get(f"/api/courses/{self.course['id']}").json()["weeks"][0]["my_progress"])

    def test_upload_rejects_empty_forged_and_non_documents(self):
        for name, data, status in [("source.pdf", b"", 422), ("source.pdf", b"not a pdf", 422), ("source.docx", b"not docx", 422), ("script.html", b"<script>", 415)]:
            result = self.client.post("/api/syllabi/upload", data={"course_id": self.course["id"]}, files={"file": (name, data)})
            self.assertEqual(result.status_code, status, result.text)
        self.assertFalse(list(Path(self.temp.name).rglob("*.pdf")))

    def test_syllabus_review_publish_visibility_and_source_deletion(self):
        syllabus = self.upload()
        self.assertEqual(syllabus["filename"], "syllabus.docx")
        self.actor = self.student
        self.assertEqual(self.client.get("/api/syllabi").json(), [])
        self.assertEqual(self.client.get(f"/api/syllabi/{syllabus['id']}/download").status_code, 404)
        self.actor = self.admin
        publish = f"/api/syllabi/{syllabus['id']}/publish"
        self.assertEqual(self.client.post(publish, json={}).status_code, 409)
        draft = {"title": "Профессиональный русский язык", "summary": "Проверенный учебный план", "weeks": [{"week_number": 1, "title": "Профессиональное общение", "tasks": "Составить диалог"}], "works": [{"title": "Практика", "week_number": 1}]}
        self.assertEqual(self.client.put(f"/api/syllabi/{syllabus['id']}", json={"draft": draft}).status_code, 200)
        result = self.client.post(publish, json={})
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["weeks_created"], 1)
        self.assertEqual(self.client.post(publish, json={}).status_code, 409)
        self.actor = self.student
        self.assertEqual(len(self.client.get("/api/syllabi").json()), 1)
        self.assertEqual(self.client.get(f"/api/syllabi/{syllabus['id']}/download").status_code, 200)
        self.actor = self.admin
        self.assertEqual(self.client.delete(f"/api/syllabi/{syllabus['id']}").status_code, 200)
        course = self.client.get(f"/api/courses/{self.course['id']}").json()
        self.assertEqual(len(course["weeks"]), 1)
        self.assertEqual(len(course["works"]), 1)
        self.assertIsNone(course["weeks"][0]["source_syllabus_id"])
        self.assertFalse(list(Path(self.temp.name).rglob("*.docx")))

    def test_publish_conflict_is_atomic_and_replacement_preserves_answers(self):
        work = self.create_work()
        self.client.post(f"/api/courses/{self.course['id']}/weeks", json={"week_number": 1, "title": "Оригинал"})
        syllabus = self.upload()
        self.client.put(f"/api/syllabi/{syllabus['id']}", json={"draft": {"weeks": [{"week_number": 1, "title": "Новая тема"}]}})
        endpoint = f"/api/syllabi/{syllabus['id']}/publish"
        self.assertEqual(self.client.post(endpoint, json={}).status_code, 409)
        self.actor = self.student
        self.client.put(f"/api/works/{work['id']}/submission", json={"content": "Ответ"})
        self.actor = self.admin
        self.assertEqual(self.client.post(endpoint, json={"replace_existing": True}).status_code, 409)
        self.assertEqual(self.client.get(f"/api/courses/{self.course['id']}").json()["weeks"][0]["title"], "Оригинал")

    def test_ai_analysis_is_queued_and_automatically_publishes_empty_course(self):
        syllabus = self.upload()
        with patch.object(settings, "GEMINI_API_KEY", ""):
            result = self.client.post(f"/api/syllabi/{syllabus['id']}/analyze")
            self.assertEqual(result.json()["status"], "needs_configuration")
        with patch.object(settings, "GEMINI_API_KEY", "test-key"), patch.object(background, "SessionLocal", self.sessions), patch("app.services.syllabus_import.analyze_syllabus_text", AsyncMock(return_value=SyllabusDraft(weeks=[{"week_number": 1, "title": "Введение"}]))):
            result = self.client.post(f"/api/syllabi/{syllabus['id']}/analyze")
            self.assertEqual(result.json()["status"], "queued")
            self.assertTrue(asyncio.run(background.process_one_job()))
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(self.client.get(f"/api/syllabi/{syllabus['id']}").json()["status"], "published")
        self.assertEqual(len(self.client.get(f"/api/courses/{self.course['id']}").json()["weeks"]), 1)

    def test_untrusted_ai_json_cannot_create_records_or_change_account_fields(self):
        syllabus = self.upload()
        payload = {"title": "План", "role": "admin", "weeks": [{"week_number": 1, "title": "Тема"}]}
        with patch.object(settings, "GEMINI_API_KEY", "test-key"), patch.object(background, "SessionLocal", self.sessions), patch("app.services.syllabus.chat_completion", AsyncMock(return_value=json.dumps(payload))):
            result = self.client.post(f"/api/syllabi/{syllabus['id']}/analyze")
            self.assertTrue(asyncio.run(background.process_one_job()))
        self.assertEqual(result.status_code, 200, result.text)
        detail = self.client.get(f"/api/syllabi/{syllabus['id']}").json()
        self.assertEqual(detail["status"], "import_error")
        self.assertEqual(self.client.get(f"/api/courses/{self.course['id']}").json()["weeks"], [])
        self.assertEqual(self.client.put(f"/api/syllabi/{syllabus['id']}", json={"draft": payload}).status_code, 422)
        with self.sessions() as db:
            self.assertEqual(db.get(User, self.student.id).role, "student")

    def test_group_deletion_cleans_children_and_preserves_accounts(self):
        work = self.create_work()
        self.upload()
        self.actor = self.student
        self.client.put(f"/api/works/{work['id']}/submission", json={"content": "Личная работа"})
        self.actor = self.admin
        result = self.client.delete(f"/api/groups/{self.group_id}")
        self.assertEqual(result.status_code, 200, result.text)
        with self.sessions() as db:
            for model in [Course, ModuleWork, WorkSubmission, SyllabusDocument]:
                self.assertEqual(db.query(model).count(), 0)
            self.assertEqual(db.query(User).count(), 4)
            self.assertIsNone(db.get(User, self.student.id).group_id)
        self.assertFalse(list(Path(self.temp.name).rglob("*.docx")))

    def test_deleted_default_group_is_not_recreated_after_restart(self):
        self.client.delete(f"/api/groups/{self.group_id}")
        with self.sessions() as db:
            seed_academic_defaults(db)
            self.assertIsNone(db.query(StudentGroup).filter_by(name="ИП-26-3кс").first())

    def test_reviewed_source_mapping_has_all_weeks_and_module_controls(self):
        mapping = json.loads(REFERENCE_DATA.read_text(encoding="utf-8"))
        draft = SyllabusDraft.model_validate(mapping["draft"])
        self.assertEqual([week.week_number for week in draft.weeks], list(range(1, 16)))
        self.assertEqual(len(draft.works), 8)
        self.assertEqual([work.week_number for work in draft.works if work.kind == "module"], [7, 15])
        self.assertEqual(mapping["course"]["credits"], 4)
        self.assertIn("Негізгі нормативтік құжаттарға шолу", draft.weeks[0].title)
        self.assertIn("Мобильді ақпараттық технологиялар", draft.weeks[-1].title)

    def test_source_import_is_idempotent_preserves_edits_and_restricts_original(self):
        mapping = json.loads(REFERENCE_DATA.read_text(encoding="utf-8"))
        source = Path(__file__).resolve().parents[2] / "local-documents" / mapping["source_filename"]
        if not source.is_file():
            self.skipTest("Provided source DOCX is not present in this checkout")
        original = source.read_bytes()
        with self.sessions() as db:
            result = seed_sourced_curriculum(db, source_path=source, upload_dir=self.temp.name)
            self.assertTrue(result["imported"])
            course_id = result["course_id"]
            course = db.get(Course, course_id)
            course.title = "Название от администратора"
            db.commit()
            self.assertFalse(seed_sourced_curriculum(db, source_path=source, upload_dir=self.temp.name)["imported"])
            self.assertEqual(db.get(Course, course_id).title, "Название от администратора")
            document = db.query(SyllabusDocument).filter_by(course_id=course_id).one()
            document_id = document.id
            self.assertTrue(document.source_restricted)
            self.assertNotIn("пароль", document.extracted_text.lower())
            self.assertEqual(db.query(CurriculumWeek).filter_by(course_id=course_id).count(), 15)
            self.assertEqual(db.query(ModuleWork).filter_by(course_id=course_id).count(), 8)
            self.assertTrue(all(work.due_date is None for work in db.query(ModuleWork).filter_by(course_id=course_id)))
        self.assertEqual(self.client.get(f"/api/groups/{self.group_id}").json()["schedule"], [])
        self.actor = self.student
        self.assertEqual(self.client.get(f"/api/syllabi/{document_id}/download").status_code, 403)
        detail = self.client.get(f"/api/syllabi/{document_id}").json()
        self.assertIn("Данные доступа исключены", detail["extracted_text"])
        self.assertEqual(source.read_bytes(), original)
        self.actor = self.admin
        self.assertEqual(self.client.get(f"/api/syllabi/{document_id}/download").status_code, 200)

    def test_syllabus_password_lines_are_not_sent_to_ai(self):
        import asyncio
        with patch("app.services.syllabus.chat_completion", AsyncMock(return_value='{"weeks": [], "works": []}')) as completion:
            asyncio.run(analyze_syllabus_text("Учебный план\nMeeting пароль ACCESS-ONLY-FIXTURE\nНеделя 1: Введение"))
        outbound = completion.call_args.args[1][0]["parts"][0]["text"]
        self.assertNotIn("ACCESS-ONLY-FIXTURE", outbound)
        self.assertIn("Неделя 1", outbound)

    def test_chat_uses_published_group_curriculum_without_private_or_draft_data(self):
        work = self.create_work()
        with self.sessions() as db:
            documents = [SyllabusDocument(course_id=self.course["id"], title="Опубликованный", filename="source.docx", filepath="unused", size=10,
                                           status="published", extracted_text="Опубликованный учебный план\nпароль NEVER-SEND-CREDENTIAL\nИскать доказательства по учебному плану."),
                         SyllabusDocument(course_id=self.course["id"], title="Секретный черновик", filename="draft.docx", filepath="unused", size=10,
                                           status="review", extracted_text="DRAFT-PRIVATE-TOPIC")]
            db.add_all(documents)
            db.flush()
            db.add(CurriculumWeek(course_id=self.course["id"], week_number=7, title="Контрольный коллоквиум", tasks="Подготовить OWN-GROUP-TASK", source_syllabus_id=documents[0].id))
            other_course = Course(group_id=self.other_group_id, title="OTHER-GROUP-COURSE")
            db.add(other_course)
            db.flush()
            other_course_id = other_course.id
            db.add(CurriculumWeek(course_id=other_course_id, week_number=7, title="OTHER-GROUP-WEEK"))
            db.add(WorkSubmission(work_id=work["id"], user_id=self.peer.id, content="PEER-PRIVATE-ANSWER", status="submitted"))
            db.commit()
        self.actor = self.student
        with patch("app.routers.chat.chat_completion", AsyncMock(return_value="Ответ по опубликованному силабусу")) as completion:
            result = self.client.post("/api/chat/send", json={"message": "Что делать на 7 неделе?", "context": {"course_id": self.course["id"], "week_number": 7}})
        self.assertEqual(result.status_code, 200, result.text)
        system = completion.call_args.args[0]
        self.assertIn("OWN-GROUP-TASK", system)
        self.assertIn("source.docx", system)
        for forbidden in ["NEVER-SEND-CREDENTIAL", "DRAFT-PRIVATE-TOPIC", "OTHER-GROUP-COURSE", "OTHER-GROUP-WEEK", "PEER-PRIVATE-ANSWER"]:
            self.assertNotIn(forbidden, system)
        with self.sessions() as db:
            count = db.query(ChatSession).count()
            self.assertLessEqual(len(academic_context(db, self.student, "7 апта")), MAX_ACADEMIC_CONTEXT)
        rejected = self.client.post("/api/chat/send", json={"message": "Покажи план", "context": {"course_id": other_course_id}})
        self.assertEqual(rejected.status_code, 403)
        with self.sessions() as db:
            self.assertEqual(db.query(ChatSession).count(), count)

    def test_chat_rejects_invalid_course_context_and_handles_missing_group(self):
        self.actor = self.student
        for context in [{"course_id": True}, {"course_id": -1}, {"course_id": "1"}, {"week_number": 53}]:
            self.assertEqual(self.client.post("/api/chat/send", json={"message": "План", "context": context}).status_code, 422)
        with self.sessions() as db:
            self.student.group_id = None
            self.assertEqual(academic_context(db, self.student, "План"), "")


if __name__ == "__main__":
    unittest.main()
