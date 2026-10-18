"""Opt-in QTests integration through the compiled production Node BFF.

Run with --run after building frontend. Disposable loopback servers and a fresh
SQLite fixture are used; live accounts, .env, uploads and paid AI are untouched.
"""
from __future__ import annotations

import argparse
from contextlib import ExitStack, closing
import io
import json
from pathlib import Path
import secrets
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time

from lumaball_stack import ROOT, StackFailure, base_environment, free_port, require, stop_child, wait_healthy


PASSWORD = "FixtureQTests4729"
FIXTURE = '''import os
from contextlib import asynccontextmanager
from fastapi import FastAPI
from app.database import Base, engine, SessionLocal
from app.models import User, Note, TestSession
from app.auth import hash_password
from app.routers import auth_routes, qtests, branding
from app.services import background

@asynccontextmanager
async def lifespan(app):
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        if db.get(User, 1) is None:
            db.add_all([User(id=1, email='qtests-admin@example.com', name='Administrator', role='admin', hashed_password=hash_password('FixtureQTests4729')),
                User(id=2, email='qtests-student@example.com', name='Student', role='student', hashed_password=hash_password('FixtureQTests4729')),
                User(id=3, email='qtests-other@example.com', name='Other student', role='student', hashed_password=hash_password('FixtureQTests4729'))])
            db.flush()
            db.add(Note(id=1, user_id=2, title='Preserved existing note', content='Existing learning content'))
            db.add(TestSession(id=1, user_id=2, status='finished', score=75, summary='Preserved earlier academic result'))
            db.commit()
    if os.environ.get('QTESTS_FIXTURE_WORKERS') == '1':
        await background.start_workers()
    try:
        yield
    finally:
        await background.stop_workers()
        engine.dispose()

app = FastAPI(lifespan=lifespan)
for module in (auth_routes, qtests, branding):
    app.include_router(module.router, prefix='/api')

@app.get('/api/health')
def health():
    return {'service':'QStudy', 'status':'ok'}
'''


def checked(response, status=200):
    require(response.status_code == status, f"HTTP {response.status_code}, expected {status}")
    return response.json()


def login(client, email):
    response = client.post("/api/auth/login", data={"username": email, "password": PASSWORD})
    value = checked(response)
    require("access_token" not in value and client.cookies.get("qstudy_session"), "Q-ID login did not create its private session cookie")
    cookie = response.headers.get("set-cookie", "").lower()
    require("httponly" in cookie and "samesite=strict" in cookie, "Q-ID cookie protections were lost")


def document_bank():
    from docx import Document
    document = Document()
    document.add_paragraph("A controlled 300-question import fixture")
    for number in range(1, 301):
        document.add_paragraph(f"{number}. Fixture question {number}?")
        document.add_paragraph(f"A) Correct answer {number}")
        document.add_paragraph(f"B) Wrong answer B {number}")
        document.add_paragraph(f"C) Wrong answer C {number}")
        document.add_paragraph(f"D) Wrong answer D {number}")
        document.add_paragraph(f"Explanation: Source explanation for question {number}: this is the documented correct answer.")
    target = io.BytesIO()
    document.save(target)
    return target.getvalue()


def ready_bank(client, bank_id):
    deadline = time.monotonic() + 25
    while time.monotonic() < deadline:
        bank = checked(client.get(f"/api/qtests/banks/{bank_id}"))
        if bank["status"] == "ready":
            return bank
        require(bank["status"] != "error", "Durable Word import failed")
        time.sleep(0.1)
    raise StackFailure("Durable Word import did not finish before its deadline")


def answer_index(question, *, correct):
    return next(index for index, text in enumerate(question["options"])
                if text.startswith("Correct answer ") == correct)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", action="store_true")
    args = parser.parse_args()
    if not args.run:
        print("SKIP: pass --run after building frontend.")
        return 0
    import httpx
    stage = "fixture setup"
    try:
        require((ROOT / "frontend/.next/BUILD_ID").is_file(), "Build frontend before integration")
        require(not any((ROOT / "frontend" / name).exists() for name in (".env", ".env.local", ".env.production", ".env.production.local")),
                "Isolated integration cannot load private frontend dotenv files")
        node = shutil.which("node")
        require(node is not None, "Node executable missing")
        with tempfile.TemporaryDirectory(prefix="qstudy-qtests-stack-") as temporary, ExitStack() as stack:
            folder = Path(temporary)
            (folder / "qtests_fixture.py").write_text(FIXTURE, encoding="utf-8")
            database = folder / "fixture.sqlite3"
            api_port = free_port()
            site_port = free_port((api_port,))
            api_url, site_url = f"http://127.0.0.1:{api_port}", f"http://127.0.0.1:{site_port}"
            environment = base_environment()
            environment.update({"SECRET_KEY": secrets.token_hex(32), "BFF_PROXY_SECRET": secrets.token_hex(32),
                "DATABASE_URL": "sqlite:///" + database.as_posix(), "UPLOAD_DIR": str(folder / "uploads"),
                "BACKUP_DIR": str(folder / "backups"), "BACKUPS_ENABLED": "false", "RATE_LIMIT_ENABLED": "false",
                "QSTUDY_MANAGED_RUNTIME": "1", "PUBLIC_BASE_URL": site_url, "CORS_ORIGINS": site_url,
                "ALLOW_PUBLIC_HTTP": "false", "RECOVERY_DELIVERY": "disabled", "GEMINI_API_KEY": "",
                "SMTP_HOST": "", "SMTP_PASSWORD": "", "PYTHONPATH": str(ROOT / "backend"), "PYTHONUTF8": "1",
                "PYTHONUNBUFFERED": "1", "NODE_ENV": "production", "NEXT_TELEMETRY_DISABLED": "1",
                "TRUST_PROXY_HEADERS": "true", "QSTUDY_ACCESS_MODE": "local", "BACKEND_URL": api_url,
                "QTESTS_FIXTURE_WORKERS": "0", "JOB_WORKERS": "1"})
            api_log = stack.enter_context((folder / "api.log").open("wb"))
            frontend_log = stack.enter_context((folder / "frontend.log").open("wb"))
            api, frontend = None, None

            def start_api():
                child = subprocess.Popen([sys.executable, "-m", "uvicorn", "qtests_fixture:app", "--host", "127.0.0.1", "--port", str(api_port),
                                          "--no-access-log", "--log-level", "error"],
                    cwd=folder, env=environment, stdin=subprocess.DEVNULL, stdout=api_log, stderr=subprocess.STDOUT)
                try:
                    wait_healthy(api_url + "/api/health", child, seconds=35, log_path=folder / "api.log")
                except Exception:
                    stop_child(child)
                    raise
                return child

            try:
                stage = "isolated API and production BFF startup"
                api = start_api()
                frontend = subprocess.Popen([node, str(ROOT / "frontend/server.mjs"), "--hostname", "127.0.0.1", "--port", str(site_port)],
                    cwd=ROOT / "frontend", env=environment, stdin=subprocess.DEVNULL, stdout=frontend_log, stderr=subprocess.STDOUT)
                wait_healthy(site_url + "/api/health", frontend, log_path=folder / "frontend.log")

                def client(origin=None):
                    return stack.enter_context(httpx.Client(base_url=site_url, headers={"Origin": origin or site_url}, timeout=25, trust_env=False))

                public, admin, student, other = client(), client(), client(), client()
                stage = "public landing, common login and QControl role boundaries"
                require(public.get("/qtests").status_code == 200, "QTests landing did not render")
                require(public.get("/brand/qtests-prototype.svg").status_code == 200, "QTests prototype logo missing")
                require(not checked(public.get("/api/qtests/config"))["ai_available"], "Fixture tried to use paid AI")
                require(public.get("/api/qtests/overview").status_code == 401, "Anonymous source banks were exposed")
                for account, email in ((admin, "qtests-admin@example.com"), (student, "qtests-student@example.com"), (other, "qtests-other@example.com")):
                    login(account, email)
                require(checked(student.get("/api/auth/me"))["email"] == "qtests-student@example.com", "QTests did not share the Q-ID account")
                require(student.get("/qtests/library").status_code == 200, "Authenticated library did not render")
                restricted = public.get("/qcontrol/qtests")
                require(restricted.status_code == 307 and restricted.headers.get("location") == "/login?next=%2Fqcontrol", "Anonymous QControl access did not return to Q-ID")
                restricted = student.get("/qcontrol/qtests")
                require(restricted.status_code == 307 and restricted.headers.get("location") == "/dashboard", "Student received QControl page")
                require(admin.get("/qcontrol/qtests").status_code == 200, "Administrator QTests panel did not render")
                require(student.get("/api/admin/qtests/overview").status_code == 403, "Student received global QTests management")

                stage = "CSRF and upload validation"
                cross_site = client("https://untrusted.example")
                cross_site.cookies.update(student.cookies)
                require(cross_site.post("/api/qtests/banks/manual", json={"title": "Forged"}).status_code == 403, "BFF accepted a cross-site mutation")
                require(student.post("/api/qtests/banks", files={"file": ("old.doc", b"legacy", "application/msword")}, data={"answer_mode": "first"}).status_code == 422,
                        "Old Word binary was silently treated as DOCX")
                require(student.post("/api/qtests/banks", files={"file": ("broken.docx", b"not-docx", "application/octet-stream")}, data={"answer_mode": "first"}).status_code == 422,
                        "Invalid DOCX content accepted")

                stage = "300-question upload persists before a worker starts"
                uploaded = checked(student.post("/api/qtests/banks", files={"file": ("fixture-300.docx", document_bank(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                    data={"title": "Controlled 300-question bank", "answer_mode": "first"}), 202)
                bank_id, job_id = uploaded["bank"]["id"], uploaded["job_id"]
                require(uploaded["bank"]["status"] == "importing" and job_id, "Upload was not queued durably")
                require(other.get(f"/api/qtests/banks/{bank_id}").status_code == 404, "Other student read the private source bank")
                with closing(sqlite3.connect(database)) as db:
                    require(db.execute("SELECT status FROM ai_jobs WHERE id=?", (job_id,)).fetchone() == ("queued",), "Import was not persisted in the job queue")
                    require(db.execute("SELECT length(source_payload) FROM qtests_banks WHERE id=?", (bank_id,)).fetchone()[0] > 1000, "Word source was not preserved for restart")
                stop_child(api); api = None
                environment["QTESTS_FIXTURE_WORKERS"] = "1"
                api = start_api()
                bank = ready_bank(student, bank_id)
                require(bank["question_count"] == 300 and len(bank["questions"]) == 300, "Restart import truncated the source bank")
                require(all(question["correct_index"] == 0 for question in bank["questions"]), "All-first import lost the source answer key")
                require(bank["warnings"], "Explicit all-first rule was not recorded")
                with closing(sqlite3.connect(database)) as db:
                    require(db.execute("SELECT status FROM ai_jobs WHERE id=?", (job_id,)).fetchone() == ("done",), "Recovered import job did not finish")
                require(checked(other.get("/api/qtests/overview"))["banks"] == [], "Bank library exposed another student")

                stage = "random 40-question attempts preserve shuffled answer mapping"
                attempt = checked(student.post(f"/api/qtests/banks/{bank_id}/attempts", json={"count": 40}), 201)
                second = checked(student.post(f"/api/qtests/banks/{bank_id}/attempts", json={"count": 40}), 201)
                attempt_id = attempt["id"]
                require(len(attempt["questions"]) == 40 and len({q["question"] for q in attempt["questions"]}) == 40, "Attempt did not select 40 unique questions")
                require({q["question"] for q in attempt["questions"]} != {q["question"] for q in second["questions"]}, "Independent attempts used the same selection")
                require(any(int(q["question"].split()[-1].rstrip("?")) > 40 for q in attempt["questions"]), "Attempt used only the first 40 source questions")
                require(any(answer_index(q, correct=True) != 0 for q in attempt["questions"]), "Correct answers remained in their source A position")
                require(all("correct_index" not in q and "explanation" not in q for q in attempt["questions"]), "Unanswered attempt exposed its answer key")
                require(other.get(f"/api/qtests/attempts/{attempt_id}").status_code == 404, "Other student read a private attempt")
                require(admin.get(f"/api/qtests/attempts/{attempt_id}").status_code == 200, "Administrator could not review an attempt")

                stage = "wrong-answer explanation and answer immutability"
                first = attempt["questions"][0]
                mutation = {"question_id": first["id"], "answer_index": answer_index(first, correct=False)}
                require(admin.post(f"/api/qtests/attempts/{attempt_id}/answer", json=mutation).status_code == 403, "Administrator answered for a student")
                answered = checked(student.post(f"/api/qtests/attempts/{attempt_id}/answer", json=mutation))
                feedback = answered["questions"][0]
                require(feedback["correct_index"] == answer_index(first, correct=True), "Option shuffle corrupted the correct answer")
                require(feedback["user_index"] != feedback["correct_index"] and feedback["explanation"].startswith("Source explanation"), "Incorrect answer did not show its source explanation")
                require(all("correct_index" not in q for q in answered["questions"][1:]), "One answer revealed future answer keys")
                require(student.post(f"/api/qtests/attempts/{attempt_id}/answer", json=mutation).status_code == 409, "Saved answer could be replaced")

                stage = "editing the bank cannot rewrite a running attempt"
                source = next(q for q in bank["questions"] if q["question"] == first["question"])
                changed_source = {"question": "Edited source after starting an attempt", "options": ["Changed A", "Changed B"], "correct_index": 1, "explanation": "Changed explanation"}
                checked(student.patch(f"/api/qtests/banks/{bank_id}/questions/{source['id']}", json=changed_source))
                snapshot = checked(student.get(f"/api/qtests/attempts/{attempt_id}"))
                require(snapshot["questions"] == answered["questions"], "Bank editing changed an existing attempt snapshot")
                require(student.get(f"/qtests/banks/{bank_id}").status_code == 200 and student.get(f"/qtests/attempts/{attempt_id}").status_code == 200,
                        "Bank editor or attempt page did not render")

                stage = "restart keeps accounts, imported bank and partial answers"
                stop_child(api); api = None
                api = start_api()
                require(checked(student.get("/api/auth/me"))["email"] == "qtests-student@example.com", "Existing Q-ID session was lost after restart")
                restored = checked(student.get(f"/api/qtests/attempts/{attempt_id}"))
                require(restored == snapshot and restored["answered_count"] == 1, "Attempt state or shuffled options changed after restart")
                require(checked(student.get(f"/api/qtests/banks/{bank_id}"))["question_count"] == 300, "Question bank was lost after restart")
                require(checked(admin.get("/api/admin/qtests/overview"))["stats"]["questions"] == 300, "QControl did not report actual imported data")

                stage = "complete 40 questions and verify the final result"
                for question in restored["questions"][1:]:
                    checked(student.post(f"/api/qtests/attempts/{attempt_id}/answer", json={"question_id": question["id"], "answer_index": answer_index(question, correct=True)}))
                finished = checked(student.post(f"/api/qtests/attempts/{attempt_id}/finish"))
                require(finished["status"] == "finished" and finished["answered_count"] == 40 and finished["correct_count"] == 39 and finished["score"] == 98,
                        "Final result did not match 39 correct answers out of 40")
                require(student.post(f"/api/qtests/attempts/{attempt_id}/answer", json=mutation).status_code == 409, "Finished result could be modified")

                stage = "administrator settings and source deletion preserve past results"
                settings = checked(public.get("/api/qtests/config"))
                editable = {name: settings[name] for name in ("enabled", "title", "description", "announcement", "default_question_count")}
                editable["enabled"] = False
                require(student.patch("/api/admin/qtests/settings", json=editable).status_code == 403, "Student changed global settings")
                checked(admin.patch("/api/admin/qtests/settings", json=editable))
                require(student.get("/api/qtests/overview").status_code == 503, "Administrator disable switch was ignored")
                editable["enabled"] = True
                checked(admin.patch("/api/admin/qtests/settings", json=editable))
                checked(student.delete(f"/api/qtests/banks/{bank_id}"))
                preserved = checked(student.get(f"/api/qtests/attempts/{attempt_id}"))
                require(preserved["bank_id"] is None and preserved["questions"] == finished["questions"] and preserved["score"] == 98,
                        "Deleting the source bank destroyed historical results")
                with closing(sqlite3.connect(database)) as db:
                    require(db.execute("SELECT COUNT(*) FROM users").fetchone() == (3,), "QTests created separate identity accounts")
                    require(db.execute("SELECT content FROM notes WHERE id=1").fetchone() == ("Existing learning content",), "Existing QStudy note was changed")
                    require(db.execute("SELECT score FROM test_sessions WHERE id=1").fetchone() == (75,), "QTests overwrote academic test history")
                    require(db.execute("PRAGMA integrity_check").fetchone() == ("ok",) and not db.execute("PRAGMA foreign_key_check").fetchall(), "Shared database integrity failed")
                print("PASS: production HTTP/BFF QTests landing and QControl role gate, shared HttpOnly Q-ID login, CSRF, safe DOCX validation, durable 300-question import across API restart, private banks, random 40 unique questions, independently shuffled correct-answer mapping, hidden keys, source explanations, immutable answers and attempt snapshots, resumed attempts, complete 39/40 grading, administrator settings, source deletion with result preservation, and unchanged earlier QStudy data. No paid AI used.")
                return 0
            finally:
                stop_child(frontend)
                stop_child(api)
    except Exception as error:
        reason = str(error) if isinstance(error, StackFailure) else type(error).__name__
        print(f"FAIL: QTests isolated stack at {stage}: {reason}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
